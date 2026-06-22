#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Build the defense .pptx from 答辩稿-20页.md, reusing PPT.pptx's template look
(per-slide background images + red title styling)."""
import os, re, zipfile
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.oxml import parse_xml

HERE = os.path.dirname(os.path.abspath(__file__))
MD   = os.path.join(HERE, "答辩稿-20页.md")
TPL  = os.path.join(HERE, "..", "PPT.pptx")
OUT  = os.path.join(HERE, "答辩-科研实践.pptx")
MEDIA_DIR = os.path.join(HERE, "assets", "_tpl")   # extracted template backgrounds

MONO  = "Menlo"
RED   = RGBColor(0xA0, 0x00, 0x16)   # title / accent (template)
DARK  = RGBColor(0x20, 0x20, 0x20)   # body
SUB   = RGBColor(0x60, 0x60, 0x60)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT = RGBColor(0xE8, 0xDD, 0xDF)   # subtitle on dark cover
L_BLANK = 6                          # 空白 layout

# ---------------- extract template backgrounds ----------------
def extract_bg():
    os.makedirs(MEDIA_DIR, exist_ok=True)
    z = zipfile.ZipFile(TPL)
    for n in ("image1.png", "image5.png"):
        p = os.path.join(MEDIA_DIR, n)
        if not os.path.exists(p):
            open(p, "wb").write(z.read("ppt/media/" + n))
    return {"dark": os.path.join(MEDIA_DIR, "image1.png"),
            "content": os.path.join(MEDIA_DIR, "image5.png")}

# ---------------- markdown parsing ----------------
def parse(md):
    lines = md.splitlines(); slides = []; cur = None; i = 0; started = False
    while i < len(lines):
        ln = lines[i]
        if ln.startswith("### 附："):
            break
        m = re.match(r'^## 第 (\d+) 页 — (.*)$', ln)
        if m:
            started = True
            if cur: slides.append(cur)
            cur = {"num": int(m.group(1)), "title": m.group(2).strip(), "blocks": []}
            i += 1; continue
        if not started:
            i += 1; continue
        s = ln.strip()
        if not s or s == "---" or s.startswith("<div") or s.startswith("</div") or s.startswith("<!--"):
            i += 1; continue
        im = re.match(r'^!\[(.*?)\]\((.*?)\)\s*$', s)
        if im:
            cur["blocks"].append(("img", im.group(2), im.group(1))); i += 1; continue
        if s.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                row = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not re.match(r'^[:\-\s|]+$', "|".join(row)):
                    rows.append(row)
                i += 1
            cur["blocks"].append(("table", rows)); continue
        if s.startswith("# "):
            cur["blocks"].append(("h1", s[2:].strip())); i += 1; continue
        if s.startswith("### "):
            cur["blocks"].append(("h3", s[4:].strip())); i += 1; continue
        if s.startswith("> "):
            q = s[2:].strip()
            if q.startswith("备注：") or q.startswith("备注:"):
                cur["blocks"].append(("note", q[3:].strip()))
            elif not q.startswith("视觉建议") and not q.startswith("节奏建议"):
                cur["blocks"].append(("quote", q))
            i += 1; continue
        mb = re.match(r'^(\s*)-\s+(.*)$', ln)
        if mb:
            lvl = 1 if len(mb.group(1)) >= 2 else 0
            cur["blocks"].append(("bullet", lvl, mb.group(2).strip())); i += 1; continue
        cur["blocks"].append(("para", s)); i += 1; continue
    if cur: slides.append(cur)
    return slides

# ---------------- pptx helpers ----------------
def set_font(run, size=None, bold=None, color=None, mono=False):
    f = run.font
    if size is not None: f.size = Pt(size)
    if bold is not None: f.bold = bold
    if color is not None: f.color.rgb = color
    if mono:
        f.name = MONO
        rPr = run._r.get_or_add_rPr()
        for tag in ("a:latin", "a:cs"):
            el = rPr.find(qn(tag))
            if el is None:
                el = rPr.makeelement(qn(tag), {}); rPr.append(el)
            el.set("typeface", MONO)

INLINE = re.compile(r'(\*\*.*?\*\*|`.*?`)')
def _clean(s): return s.replace("**", "").replace("`", "")
def add_runs(p, text, size=None, color=DARK, bold_all=False, code_color=RED):
    for part in INLINE.split(text):
        if not part: continue
        if part.startswith("**") and part.endswith("**"):
            r = p.add_run(); set_font(r, size, True, color); r.text = _clean(part[2:-2])
        elif part.startswith("`") and part.endswith("`"):
            r = p.add_run(); set_font(r, size, bold_all, code_color, mono=True); r.text = _clean(part[1:-1])
        else:
            r = p.add_run(); set_font(r, size, bold_all, color); r.text = _clean(part)

def set_bg(slide, img):
    _, rId = slide.part.get_or_add_image_part(img)
    bg = parse_xml(
        '<p:bg xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        '<p:bgPr><a:blipFill dpi="0" rotWithShape="1"><a:blip r:embed="%s"/>'
        '<a:stretch><a:fillRect/></a:stretch></a:blipFill><a:effectLst/></p:bgPr></p:bg>' % rId)
    slide._element.find(qn('p:cSld')).insert(0, bg)

def textbox(slide, l, t, w, h, anchor=MSO_ANCHOR.TOP):
    tf = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h)).text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    for m in ("margin_left","margin_right","margin_top","margin_bottom"): setattr(tf, m, Pt(2))
    return tf

def add_title(slide, text):
    tf = textbox(slide, 1.1, 1.02, 11.2, 0.75, MSO_ANCHOR.MIDDLE)
    add_runs(tf.paragraphs[0], text, 24, RED, bold_all=True)

def add_text_block(slide, items, l, t, w, h, dark=True):
    body = DARK if dark else LIGHT
    tf = textbox(slide, l, t, w, h); first = True
    for it in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph(); first = False
        p.space_after = Pt(7); p.line_spacing = 1.12
        if it[0] == "bullet":
            lvl, txt = it[1], it[2]; p.level = lvl
            rb = p.add_run(); set_font(rb, 15 if not lvl else 13, False, RED if not lvl else SUB)
            rb.text = "•  " if not lvl else "－ "
            add_runs(p, txt, 15 if not lvl else 13, body)
        elif it[0] == "quote":
            add_runs(p, it[1], 12.5, SUB)
        else:
            add_runs(p, it[1], 15, body)
    return tf

def img_size(path):
    from PIL import Image
    with Image.open(path) as im: return im.size

def place_image(slide, path, region):
    l, t, w, h = region; iw, ih = img_size(path); r = iw / ih
    if w / h > r: nh = h; nw = h * r
    else:         nw = w; nh = w / r
    slide.shapes.add_picture(path, Inches(l + (w - nw) / 2), Inches(t + (h - nh) / 2), Inches(nw), Inches(nh))

def md_table(slide, rows, l, t, w, h):
    nr, nc = len(rows), max(len(r) for r in rows)
    gt = slide.shapes.add_table(nr, nc, Inches(l), Inches(t), Inches(w), Inches(h)).table
    for ci in range(nc): gt.columns[ci].width = Inches(w / nc)
    for ri, row in enumerate(rows):
        for ci in range(nc):
            cell = gt.cell(ri, ci)
            cell.margin_left = Pt(6); cell.margin_right = Pt(6)
            cell.margin_top = Pt(3); cell.margin_bottom = Pt(3)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            add_runs(cell.text_frame.paragraphs[0], row[ci] if ci < len(row) else "",
                     12.5, DARK if ri else WHITE, bold_all=(ri == 0), code_color=DARK if ri else WHITE)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RED if ri == 0 else (RGBColor(0xF6,0xEC,0xEE) if ri % 2 else WHITE)
    return gt

# ---------------- slide builders ----------------
CL, CT, CW, CB = 1.1, 1.95, 11.2, 6.85

def build(prs, sl, BG):
    title, blocks = sl["title"], sl["blocks"]
    imgs   = [b for b in blocks if b[0] == "img"]
    tables = [b for b in blocks if b[0] == "table"]
    texts  = [b for b in blocks if b[0] in ("bullet", "para", "quote")]
    h1s = [b[1] for b in blocks if b[0] == "h1"]
    h3s = [b[1] for b in blocks if b[0] == "h3"]
    s = prs.slides.add_slide(prs.slide_layouts[L_BLANK])

    notes = [b[1] for b in blocks if b[0] == "note"]
    if notes:
        ntf = s.notes_slide.notes_text_frame
        ntf.text = notes[0]
        for n in notes[1:]:
            ntf.add_paragraph().text = n

    if "封面" in title:
        set_bg(s, BG["dark"])
        tf = textbox(s, 1.2, 1.55, 10.9, 1.3, MSO_ANCHOR.MIDDLE)
        add_runs(tf.paragraphs[0], h1s[0] if h1s else title, 38, RED, bold_all=True)
        tf2 = textbox(s, 1.2, 3.1, 10.9, 0.8, MSO_ANCHOR.MIDDLE)
        if h3s: add_runs(tf2.paragraphs[0], h3s[0], 19, LIGHT, bold_all=True)
        if texts:
            tf3 = textbox(s, 1.2, 4.3, 10.9, 1.8)
            for k, it in enumerate(texts):
                p = tf3.paragraphs[0] if k == 0 else tf3.add_paragraph(); p.space_after = Pt(5)
                add_runs(p, it[-1], 14, LIGHT)
        return

    if "开篇" in title or title.strip().startswith("PART"):
        # plain white section divider, big red title
        tf = textbox(s, 1.2, 2.6, 10.9, 1.4, MSO_ANCHOR.MIDDLE)
        add_runs(tf.paragraphs[0], h1s[0] if h1s else title, 44, RED, bold_all=True)
        if h3s:
            tf2 = textbox(s, 1.2, 4.1, 10.9, 0.7, MSO_ANCHOR.MIDDLE)
            add_runs(tf2.paragraphs[0], h3s[0], 18, SUB)
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.22), Inches(3.95), Inches(2.4), Pt(3))
        bar.fill.solid(); bar.fill.fore_color.rgb = RED; bar.line.fill.background(); bar.shadow.inherit = False
        return

    if "Q & A" in title or "Q&A" in title:
        set_bg(s, BG["dark"])
        tf = textbox(s, 1.2, 3.0, 10.9, 1.4, MSO_ANCHOR.MIDDLE)
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        add_runs(p, "谢谢！　Q & A", 40, RED, bold_all=True)
        return

    if "目录" in title:
        tf = textbox(s, 1.2, 0.95, 10.9, 1.0)
        add_runs(tf.paragraphs[0], "目  录", 40, RED, bold_all=True)
        if tables:
            md_table(s, tables[0][1], 1.4, 2.4, 10.5, min(4.0, 0.62 * len(tables[0][1]) + 0.5))
        return

    # ---- content (image5 background) ----
    set_bg(s, BG["content"])
    add_title(s, title)
    ch = CB - CT

    if imgs and tables:
        place_image(s, imgs[0][1], (CL, CT, 5.6, ch))
        rt = CT
        if texts:
            n = sum(2 if len(it[-1]) > 28 else 1 for it in texts)
            th = min(2.2, 0.45 + 0.42 * n); add_text_block(s, texts, 6.85, rt, 5.45, th); rt += th + 0.1
        md_table(s, tables[0][1], 6.85, rt, 5.45, CB - rt)
        return
    if imgs and texts:
        n = sum(2 if len(it[-1]) > 34 else 1 for it in texts)
        th = min(2.5, 0.4 + 0.4 * n); ih = ch - th - 0.12
        if len(imgs) == 1: place_image(s, imgs[0][1], (CL, CT, CW, ih))
        else:
            ew = CW / len(imgs)
            for k, im in enumerate(imgs): place_image(s, im[1], (CL + k * ew, CT, ew, ih))
        add_text_block(s, texts, CL, CT + ih + 0.12, CW, th)
        return
    if imgs:
        if len(imgs) == 1: place_image(s, imgs[0][1], (CL, CT, CW, ch))
        else:
            ew = CW / len(imgs)
            for k, im in enumerate(imgs): place_image(s, im[1], (CL + k * ew, CT, ew, ch))
        return

    if "演示" in title:
        ph = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.4), Inches(CT), Inches(4.9), Inches(4.3))
        ph.fill.solid(); ph.fill.fore_color.rgb = RGBColor(0x2A, 0x2A, 0x2A)
        ph.line.color.rgb = RED; ph.line.width = Pt(1.5); ph.shadow.inherit = False
        pp = ph.text_frame.paragraphs[0]; pp.alignment = PP_ALIGN.CENTER
        add_runs(pp, "▶  演示视频\n（在此插入）", 20, WHITE, bold_all=True, code_color=WHITE)
        add_text_block(s, [b for b in blocks if b[0] in ("para","bullet","quote")], CL, CT, 6.1, ch)
        return

    rt = [CT]; buf = []
    def flush():
        if not buf: return
        n = sum(2 if len(it[-1]) > 40 else 1 for it in buf)
        th = min(CB - rt[0], 0.35 + 0.42 * n)
        add_text_block(s, list(buf), CL, rt[0], CW, th); rt[0] += th + 0.12; buf.clear()
    for b in blocks:
        if b[0] in ("para","bullet","quote"): buf.append(b)
        elif b[0] == "table":
            flush(); th = min(CB - rt[0], 0.52 * len(b[1]) + 0.45)
            md_table(s, b[1], CL, rt[0], CW, th); rt[0] += th + 0.12
    flush()

def main():
    BG = extract_bg()
    prs = Presentation(TPL)
    sldIdLst = prs.slides._sldIdLst
    for sldId in list(sldIdLst):
        prs.part.drop_rel(sldId.get(qn('r:id'))); sldIdLst.remove(sldId)
    for sl in parse(open(MD, encoding="utf-8").read()):
        build(prs, sl, BG)
    prs.save(OUT)
    print(f"wrote {OUT}  ({len(prs.slides._sldIdLst)} slides; template backgrounds applied)")

if __name__ == "__main__":
    main()
